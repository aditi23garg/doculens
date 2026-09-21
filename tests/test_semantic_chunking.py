"""Offline tests for structure parsing, semantic chunking, citations and query planning."""
import io
import re
import zlib

from rag.citations import cited_numbers
from rag.query_planner import merge_queries, parse_queries
from rag.semantic_chunker import semantic_chunks
from rag.structure import parse_sections


def fake_embed(texts):
    """Deterministic bag-of-words vectors: same vocabulary => similar, different => distant."""
    out = []
    for t in texts:
        v = [0.0] * 256
        for w in re.findall(r"[a-z]+", t.lower()):
            v[zlib.crc32(w.encode()) % 256] += 1.0
        out.append(v)
    return out


def _pages(text):
    return [(None, text)]


# ── structure ─────────────────────────────────────────────────────────────────
def test_markdown_headings_build_a_path():
    text = "# Phase 1\n\nIntro text.\n\n## PaddleOCR\n\nDetails about paddle.\n\n## GLM-OCR\n\nDetails about glm."
    secs = parse_sections(_pages(text))
    assert [s.path for s in secs] == [("Phase 1",), ("Phase 1", "PaddleOCR"), ("Phase 1", "GLM-OCR")]


def test_heuristic_headings_only_when_enabled():
    text = "2.1 Setup steps\n\nInstall the package first.\n\nRESULTS SUMMARY\n\nAll tests passed."
    assert parse_sections(_pages(text))[0].path == ()  # off: one anonymous section
    paths = [s.path for s in parse_sections(_pages(text), heuristic_headings=True)]
    assert paths == [("2.1 Setup steps",), ("RESULTS SUMMARY",)]


def test_ordinary_sentences_are_not_headings():
    text = "This is a normal sentence that ends with a period.\n\nAnother line, with a comma,"
    assert [s.path for s in parse_sections(_pages(text), heuristic_headings=True)] == [()]


# ── semantic chunking ─────────────────────────────────────────────────────────
def test_small_section_stays_whole_with_its_heading_path():
    secs = parse_sections(_pages("# Models\n\n## GLM-OCR\n\nGLM-OCR is a compact model. It reads scanned forms."))
    chunks = semantic_chunks(secs, fake_embed, max_chars=1800, min_chars=10)
    assert len(chunks) == 1 and chunks[0].section == "Models > GLM-OCR" and "compact model" in chunks[0].text


def test_each_model_section_becomes_its_own_chunk():
    body = " ".join(f"Sentence {i} explains the details of this model." for i in range(8))
    text = "\n\n".join(f"## Model {n}\n\n{body}" for n in "ABC")
    chunks = semantic_chunks(parse_sections(_pages(text)), fake_embed, max_chars=1800, min_chars=100)
    assert [c.section for c in chunks] == ["Model A", "Model B", "Model C"]


def test_tiny_neighbouring_sections_are_merged():
    text = "## One\n\nShort text.\n\n## Two\n\nAlso short.\n\n## Three\n\nStill short."
    chunks = semantic_chunks(parse_sections(_pages(text)), fake_embed, max_chars=1800, min_chars=300)
    assert len(chunks) == 1
    assert "Two" in chunks[0].text and "Three" in chunks[0].text  # headings kept inline when merged


def test_long_section_splits_at_the_topic_shift():
    a = [f"The paddle vision model reads printed text and handwriting number {i}." for i in range(14)]
    b = [f"The insurance invoice lists premium payment and claim billing amount {i}." for i in range(14)]
    text = "## Mixed\n\n" + " ".join(a + b)
    chunks = semantic_chunks(parse_sections(_pages(text)), fake_embed, max_chars=1400, min_chars=100)
    assert len(chunks) >= 2
    for c in chunks:  # no chunk mixes both topics
        assert not ("paddle" in c.text and "insurance" in c.text), c.text
    joined = " ".join(c.text for c in chunks)
    assert all(sentence in joined for sentence in a + b)  # nothing lost


def test_embedding_failure_falls_back_to_size_packing():
    def broken(_):
        raise RuntimeError("rate limited")

    text = "## Long\n\n" + " ".join(f"Sentence number {i} is here." for i in range(120))
    chunks = semantic_chunks(parse_sections(_pages(text)), broken, max_chars=600, min_chars=100)
    assert len(chunks) > 3 and all(len(c.text) <= 600 for c in chunks)


def test_page_range_is_tracked():
    pages = [(3, "## Costs\n\nFirst part."), (4, "Second part continues here.")]
    chunks = semantic_chunks(parse_sections(pages), fake_embed, min_chars=10)
    assert (chunks[0].page, chunks[0].page_end) == (3, 4)


# ── citations ─────────────────────────────────────────────────────────────────
def test_all_bracket_styles_are_recognised():
    assert cited_numbers("A [1]. B 【2】. C [3, 4]. D 〔5〕.") == {1, 2, 3, 4, 5}


def test_years_are_not_citations():
    assert cited_numbers("Released in [2023] and again [2024].") == set()


# ── query planning ────────────────────────────────────────────────────────────
def test_parse_queries_tolerates_chatter_and_garbage():
    assert parse_queries('Sure! ["models compared", "candidate models"]') == ["models compared", "candidate models"]
    assert parse_queries("no json here") == []
    assert parse_queries('["ok", 5, "", "fine"]') == ["ok", "fine"]


def test_merge_queries_dedupes_and_caps():
    assert merge_queries("Q", ["q", "a", "b", "c"], limit=3) == ["Q", "a", "b"]


# ── DOCX extraction keeps headings and order (uses real python-docx) ──────────
def test_docx_headings_survive_extraction_in_order():
    import sys, types
    for name in ("config.settings",):
        sys.modules.setdefault(name, types.ModuleType(name))
    sys.modules["config.settings"].get_settings = lambda: None
    for name in ("database.mongo_client", "database.vector_store", "rag.embeddings"):
        sys.modules.setdefault(name, types.ModuleType(name))
    import database
    database.mongo_client = sys.modules["database.mongo_client"]
    database.vector_store = sys.modules["database.vector_store"]
    sys.modules["rag.embeddings"].embed_documents = lambda *a, **k: []
    import docx
    from rag.document_processor import extract_pages

    d = docx.Document()
    d.add_heading("OCR Model Research", level=1)
    d.add_paragraph("Why we are doing this.")
    d.add_heading("PaddleOCR-VL 1.6", level=2)
    d.add_paragraph("Paddle details.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Accuracy", "94%"
    d.add_heading("GLM-OCR", level=2)
    d.add_paragraph("GLM details.")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_pages(".docx", buf.getvalue())[0][1]
    assert text.index("# OCR Model Research") < text.index("## PaddleOCR-VL 1.6") < text.index("Accuracy | 94%") < text.index("## GLM-OCR")
    paths = [s.path for s in parse_sections([(None, text)])]
    assert paths == [("OCR Model Research",), ("OCR Model Research", "PaddleOCR-VL 1.6"), ("OCR Model Research", "GLM-OCR")]
